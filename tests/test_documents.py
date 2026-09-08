from pathlib import Path
import time,io,zipfile
from fastapi.testclient import TestClient
from backend.app import create_app
from backend.documents import retrieve,extract
from backend.exports import build_exports
from backend.workflows import extractive_approval_note
from tests.test_auth import client,register
from openpyxl import load_workbook

class Engine:
    calls=[]
    def close(self):pass
    def status(self):return []
    def generate(self,prompt,**kwargs):
        self.calls.append(kwargs)
        return {'answer':'The identification label is damaged [S1]. Costs are not supplied [S1].','model_name':'Test fixture','capability':'general','routing_reason':'Test fixture','duration_ms':1,'finish_reason':'stop'}

def wait_run(c,w):
    deadline=time.monotonic()+15
    while time.monotonic()<deadline:
        rows=c.get(f'/api/workspaces/{w}/runs').json()['runs']
        if rows and rows[0]['status']!='running':return rows[0]
        time.sleep(.1)
    raise AssertionError('Run did not complete')

def test_read_grounding_workflow_and_artifact_isolation(tmp_path):
    engine=Engine();app=create_app(tmp_path/'db.sqlite3',engine)
    with client(app) as a,client(app) as b:
        register(a);register(b,'bob')
        w=a.post('/api/workspaces',json={'name':'Source review'}).json()['id']
        doc=a.post(f'/api/workspaces/{w}/documents',params={'filename':'inspection.txt'},content=b'Equipment P-101 identification label damaged. Costs not provided.').json()['id']
        assert b.post(f'/api/documents/{doc}/extract').status_code==404
        assert a.post(f'/api/documents/{doc}/extract').status_code==200
        pages=a.get(f'/api/documents/{doc}/pages').json()['pages']
        assert pages[0]['page']==1 and 'P-101' in pages[0]['text']
        assert b.get(f'/api/documents/{doc}/pages').status_code==404
        q=a.post(f'/api/workspaces/{w}/questions',json={'prompt':'What happened to P-101?'})
        assert q.status_code==201 and q.json()['sources'][0]['document_id']==doc
        assert q.json()['citation_format_valid']
        assert 'context' in engine.calls[-1]
        assert b.post(f'/api/workspaces/{w}/questions',json={'prompt':'P-101'}).status_code==404
        absent=a.post(f'/api/workspaces/{w}/questions',json={'prompt':'UnrelatedXYZ'}).json()
        assert absent['sources']==[] and absent['capability']=='retrieval'
        run=a.post(f'/api/workspaces/{w}/runs',json={'document_id':doc})
        assert run.status_code==202
        result=wait_run(a,w)
        assert result['status']=='complete',result
        assert len(result['result']['artifacts'])==3
        for art in result['result']['artifacts']:
            response=a.get(f"/api/artifacts/{art['id']}/download")
            assert response.status_code==200 and zipfile.is_zipfile(io.BytesIO(response.content))
            assert b.get(f"/api/artifacts/{art['id']}/download").status_code==404
        assert b.get(f'/api/workspaces/{w}/runs').status_code==404

def test_export_formula_injection_and_open_formats(tmp_path):
    evidence=[{'ref':'S1','filename':'=DANGEROUS()','page':1,'text':'+payload','document_id':'test'}]
    files=build_exports(tmp_path,'Draft [S1]',evidence)
    wb=load_workbook(tmp_path/'findings-register.xlsx')
    assert wb.active['B2'].data_type=='s' and wb.active['D2'].data_type=='s'
    from docx import Document
    from pptx import Presentation
    assert any('Draft' in p.text for p in Document(tmp_path/'approval-note.docx').paragraphs)
    assert len(Presentation(tmp_path/'inspection-briefing.pptx').slides)>=3
    assert len(files)==3

def test_ocr_fixture(tmp_path):
    pages=extract(Path('samples/inspection-scan.pdf'),'scan.pdf',tmp_path,'scan')
    assert pages[0]['method']=='ocr'
    assert 'P-101' in pages[0]['text'] and 'label' in pages[0]['text']
    assert (tmp_path/'scan-1.png').exists()

def test_extractive_approval_note_is_structured():
    evidence=[{'ref':'S1','filename':'inspection.pdf','page':1,'document_id':'d1','text':'Inspection reference: DEMO-002\nEquipment: P-101\nFinding: identification label damaged\nRecommendation: replace the damaged label after review.\nCost: not provided.'}]
    note=extractive_approval_note(evidence)
    for section in ('Subject:','Background','Recorded finding','Proposed action','Financial position','Decision requested','Source references'):
        assert section in note
    assert 'P-101' in note and '[S1]' in note and 'does not authorize' in note
