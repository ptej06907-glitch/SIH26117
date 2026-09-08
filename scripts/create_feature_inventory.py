from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Aegis_Feature_Inventory.docx"

FEATURES = [
    ("Local identity", "Local account creation, secure password hashing, expiring sessions, request protection, sign-out and authentication rate limiting.", "Complete"),
    ("Private workspaces", "Per-user workspaces with ownership checks for materials, answers, workflows and deliverables.", "Complete"),
    ("Dashboard", "Overview metrics, activity, desktop and mobile navigation, onboarding guide and compact display mode.", "Complete"),
    ("Material upload", "PDF, TXT, CSV, PNG and JPEG by picker or drag and drop, with per-file and workspace limits.", "Complete"),
    ("Material explorer", "Filename search, readiness filters, downloads, file-type labels and local-storage status.", "Complete"),
    ("Document reading", "Text extraction, scanned-page OCR, page previews, extracted-text review and accuracy warnings.", "Complete"),
    ("Local model routing", "General, coding and vision roles selected automatically from the task, with no cloud fallback.", "Complete"),
    ("Grounded answers", "Local retrieval, page-linked sources, citation checks and an extractive fallback when model references fail.", "Complete"),
    ("Multimodal review", "Local image interpretation for images and the first page of scanned PDFs.", "Complete"),
    ("Inspection workflow", "Read, retrieve, draft, check references, retry once and create editable outputs.", "Complete"),
    ("Office deliverables", "Approval note in DOCX, findings register in XLSX and briefing in PPTX.", "Complete"),
    ("Verified utilities", "Fixed-purpose code generation, three independent cases, one repair attempt and downloadable source.", "Complete"),
    ("Execution isolation", "WebAssembly and WASI execution with time, fuel, memory and output limits and no host-directory or network grants.", "Complete"),
    ("Workflow history", "Live status, progress, event history, verification checks, filters and grouped downloads.", "Complete"),
    ("Offline visibility", "Loopback binding, Python outbound guard and application plus child-process connection observations.", "Prototype evidence"),
    ("Persistence and audit", "SQLite persistence for project records and local audit events.", "Complete"),
]

LIMITATIONS = [
    "DOCX, XLSX and PPTX are output formats and are not accepted as source uploads yet.",
    "Material and workspace deletion controls are not present.",
    "Retrieval is lexical and sized for a demonstration collection rather than a large enterprise corpus.",
    "The application database is not encrypted, and a Windows administrator can access workstation files.",
    "Password recovery, administrator provisioning, invitations and enterprise SSO are not included.",
    "Connection snapshots can miss short-lived traffic; packet-level capture is not yet complete.",
    "Generated text, OCR, vision, calculations and engineering interpretations require human review.",
    "The workstation configuration permits one active model generation request at a time.",
]


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd"))
    if node is None:
        node = OxmlElement("w:shd")
        props.append(node)
    node.set(qn("w:fill"), fill)


def borders(table):
    props = table._tbl.tblPr
    node = props.find(qn("w:tblBorders"))
    if node is None:
        node = OxmlElement("w:tblBorders")
        props.append(node)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        item = OxmlElement(f"w:{edge}")
        item.set(qn("w:val"), "single")
        item.set(qn("w:sz"), "4")
        item.set(qn("w:color"), "D9D9D9")
        node.append(item)


def cell_margin(cell, top=100, start=110, bottom=100, end=110):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        item = margins.find(qn(f"w:{edge}"))
        if item is None:
            item = OxmlElement(f"w:{edge}")
            margins.append(item)
        item.set(qn("w:w"), str(value))
        item.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    props = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    props.append(node)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for name, size in (("Title", 25), ("Heading 1", 17), ("Heading 2", 12)):
    style = styles[name]
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(12)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.left_indent = Inches(0)
    ppr = style.element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)

title = doc.add_paragraph(style="Title")
title.add_run("Aegis Prototype Feature Inventory")
title_ppr = title._p.get_or_add_pPr()
title_border = title_ppr.find(qn("w:pBdr"))
if title_border is not None:
    title_ppr.remove(title_border)
subtitle = doc.add_paragraph()
subtitle_run = subtitle.add_run("SIH 26117  |  Sovereign On Premise Agentic AI Workbench  |  8 September 2026")
subtitle_run.bold = True
subtitle_run.font.size = Pt(9.5)
subtitle_run.font.color.rgb = RGBColor(69, 88, 106)

doc.add_paragraph(
    "This document records the capabilities currently implemented in the Aegis workstation prototype. "
    "The end-to-end demonstration path is operational: a local user can upload and read material, ask a source-grounded question, run an inspection workflow, download editable Office files, verify a small utility in isolation and inspect local connection observations."
)

doc.add_heading("Current validation", level=1)
validation = doc.add_table(rows=2, cols=4)
validation.alignment = WD_TABLE_ALIGNMENT.CENTER
validation.autofit = False
validation.columns[0].width = Inches(1.75)
validation.columns[1].width = Inches(1.75)
validation.columns[2].width = Inches(1.75)
validation.columns[3].width = Inches(1.75)
for i, value in enumerate(("32", "3", "5", "3")):
    cell = validation.cell(0, i)
    cell.text = value
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].runs[0].font.size = Pt(18)
    cell.paragraphs[0].runs[0].bold = True
    shade(cell, "EAF1F8")
for i, value in enumerate(("Automated tests passed", "Local model roles", "Dashboard views", "Editable Office formats")):
    cell = validation.cell(1, i)
    cell.text = value
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraphs[0].runs[0].font.size = Pt(8.5)
for row in validation.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margin(cell, 120, 90, 120, 90)
borders(validation)

doc.add_heading("Implemented features", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
widths = (Inches(1.45), Inches(4.65), Inches(1.05))
for i, width in enumerate(widths):
    table.columns[i].width = width
headers = ("Area", "Capability", "State")
for i, header in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = header
    shade(cell, "17324A")
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    cell.paragraphs[0].runs[0].bold = True
    cell.paragraphs[0].runs[0].font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    cell_margin(cell)
set_repeat_header(table.rows[0])

for index, (area, capability, state) in enumerate(FEATURES):
    cells = table.add_row().cells
    for cell, value, width in zip(cells, (area, capability, state), widths):
        cell.text = value
        cell.width = width
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_margin(cell, 85, 100, 85, 100)
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(8.4)
    cells[0].paragraphs[0].runs[0].bold = True
    cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if index % 2:
        for cell in cells:
            shade(cell, "F3F7FA")
borders(table)

doc.add_heading("Current limitations", level=1)
doc.add_paragraph(
    "The following items define the boundary of the current demonstration and should remain visible during evaluation."
)
for item in LIMITATIONS:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.add_run(item)

doc.add_heading("Recommended demonstration sequence", level=1)
steps = [
    ("Sign in and open a private workspace", "Show that another account cannot see its records."),
    ("Upload and read the fictional inspection scan", "Open the page image and extracted OCR text together."),
    ("Ask a source-grounded question", "Open the page-linked evidence from the saved answer."),
    ("Create an inspection review pack", "Follow workflow events and download the DOCX, XLSX and PPTX outputs."),
    ("Generate and verify a utility", "Show the three independent checks and downloadable source file."),
    ("Open Offline status", "Show the loopback binding and current connection observations with the stated evidence limit."),
]
for number, (heading, detail) in enumerate(steps, 1):
    paragraph = doc.add_paragraph()
    lead = paragraph.add_run(f"{number}. {heading}. ")
    lead.bold = True
    paragraph.add_run(detail)

doc.add_paragraph()
closing = doc.add_paragraph()
closing.add_run("Prototype status: ").bold = True
closing.add_run("Ready for a controlled SIH demonstration using public or fictional material. Human review remains required for every generated conclusion and deliverable.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("Aegis  |  SIH 26117  |  Feature inventory")
run.font.size = Pt(8)
run.font.color.rgb = RGBColor(105, 120, 135)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
