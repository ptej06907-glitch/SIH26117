"""Deterministic Office files from one draft and its source evidence."""
from pathlib import Path
import textwrap
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from pptx import Presentation
from pptx.util import Inches as SlideInches, Pt as SlidePt
from pptx.dml.color import RGBColor as SlideRGB

def safe_cell(value):
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value

def build_exports(directory,draft,evidence):
    directory=Path(directory);directory.mkdir(parents=True,exist_ok=True)
    doc=Document();section=doc.sections[0];section.top_margin=section.bottom_margin=Inches(.7)
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
    for style in ['Title','Heading 1','Heading 2']:doc.styles[style].font.color.rgb=RGBColor(0,0,0)
    for style in doc.styles:
        for border in list(style.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    doc.add_heading('Inspection review approval note',0)
    doc.add_paragraph('Draft for human review. This note summarizes supplied evidence; it does not authorize expenditure or maintenance.')
    doc.add_heading('Proposed note',1)
    section_names={'Background','Recorded finding','Proposed action','Financial position','Decision requested','Source references'}
    for paragraph in draft.split('\n'):
        text=paragraph.strip()
        if not text:continue
        if text in section_names:
            doc.add_heading(text,2)
        elif text.startswith('Subject:'):
            item=doc.add_paragraph();item.add_run('Subject:').bold=True;item.add_run(text.split(':',1)[1])
        else:
            doc.add_paragraph(text)
    doc.add_page_break()
    doc.add_heading('Source evidence',1)
    for e in evidence:
        doc.add_heading(f"{e['ref']} {e['filename']} page {e['page']}",2);paragraph=doc.add_paragraph(e['text']);paragraph.paragraph_format.keep_together=True
    doc.add_heading('Required review',1);doc.add_paragraph('Verify identifiers, measurements, procedure revisions, costs and requested actions against the original documents. Missing information must be supplied by the responsible reviewer.')
    doc.save(directory/'approval-note.docx')
    wb=Workbook();ws=wb.active;ws.title='Evidence register'
    ws.append(['Source','Document','Page','Evidence excerpt','Review status'])
    for e in evidence:ws.append([safe_cell(e['ref']),safe_cell(e['filename']),e['page'],safe_cell(e['text']),'Needs human review'])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='18354D')
    for col,width in [('A',12),('B',32),('C',10),('D',85),('E',24)]:ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=120
    note=wb.create_sheet('Draft note');note.append(['Draft for review']);note.append([safe_cell(draft)]);note.column_dimensions['A'].width=110;note['A2'].alignment=Alignment(wrap_text=True,vertical='top');note.row_dimensions[2].height=300
    for sheet in wb:
        sheet.page_setup.orientation='landscape';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.print_options.horizontalCentered=True
        sheet.print_area=sheet.dimensions
    ws.print_title_rows='1:1'
    wb.save(directory/'findings-register.xlsx')
    deck=Presentation();deck.slide_width=SlideInches(13.333);deck.slide_height=SlideInches(7.5)
    def slide(title,body):
        s=deck.slides.add_slide(deck.slide_layouts[6]);bg=s.background.fill;bg.solid();bg.fore_color.rgb=SlideRGB(247,249,252)
        for text,y,height,size in [(title,.55,.9,28),(body,1.65,4.9,20),('DRAFT FOR HUMAN REVIEW',6.8,.3,10)]:
            box=s.shapes.add_textbox(SlideInches(.7),SlideInches(y),SlideInches(11.8),SlideInches(height));tf=box.text_frame;tf.word_wrap=True
            tf.text=text
            for p in tf.paragraphs:p.font.size=SlidePt(size);p.font.color.rgb=SlideRGB(23,45,65)
    slide('Inspection review','Source-grounded draft briefing\nReview the original documents before taking action.')
    for i,part in enumerate(textwrap.wrap(draft,650,replace_whitespace=False)[:5]):slide(f'Draft note {i+1}',part)
    for e in evidence:slide(f"Evidence {e['ref']}",f"{e['filename']} · Page {e['page']}\n\n{e['text'][:550]}")
    slide('Reviewer actions','Confirm the source revision and OCR transcription.\nCheck each claim against its cited source.\nSupply missing costs, measurements and authorization details.')
    deck.save(directory/'inspection-briefing.pptx')
    return ['approval-note.docx','findings-register.xlsx','inspection-briefing.pptx']
